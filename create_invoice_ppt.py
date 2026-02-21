#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import argparse
import os
import shutil
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


def create_ppt(input_root: str, persons: list, output_ppt: str) -> int:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    image_extensions = {'.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tif', '.tiff'}
    slide_count = 0

    for person_name in persons:
        folder_path = os.path.join(input_root, person_name)
        if not os.path.exists(folder_path):
            print(f"跳过不存在文件夹: {folder_path}", file=sys.stderr)
            continue

        image_files = [
            os.path.join(folder_path, f)
            for f in sorted(os.listdir(folder_path))
            if not f.startswith('.') and Path(f).suffix.lower() in image_extensions and os.path.isfile(os.path.join(folder_path, f))
        ]

        for image_file in image_files:
            blank_slide_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank_slide_layout)

            title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.5))
            title_frame = title_box.text_frame
            title_frame.text = os.path.basename(image_file)
            title_frame.paragraphs[0].font.size = Pt(18)
            title_frame.paragraphs[0].font.bold = True

            person_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(2), Inches(0.3))
            person_frame = person_box.text_frame
            person_frame.text = f"姓名: {Path(person_name).name}"
            person_frame.paragraphs[0].font.size = Pt(12)

            left = Inches(0.5)
            top = Inches(1)
            max_width = Inches(9)
            max_height = Inches(6)

            pic = slide.shapes.add_picture(image_file, left, top, width=max_width)
            if pic.height > max_height:
                aspect_ratio = pic.width / pic.height
                pic.height = int(max_height)
                pic.width = int(max_height * aspect_ratio)

            slide_count += 1
            print(f"已添加: {person_name} - {os.path.basename(image_file)}")

    prs.save(output_ppt)
    print(f"\n✓ PPT已保存: {output_ppt}")
    print(f"总共添加了 {slide_count} 张幻灯片")
    return slide_count


def convert_to_pdf(ppt_path: str, outdir: str) -> bool:
    soffice = shutil.which("soffice")
    if not soffice:
        mac_default = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
        if os.path.exists(mac_default):
            soffice = mac_default

    if not soffice:
        print("LibreOffice 未安装或 soffice 路径不可用，跳过 PDF 转换", file=sys.stderr)
        return False

    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", ppt_path, "--outdir", outdir], check=True)
        print(f"✓ 已生成 PDF: {os.path.join(outdir, Path(ppt_path).with_suffix('.pdf').name)}")
        return True
    except subprocess.CalledProcessError:
        print("PDF 转换失败", file=sys.stderr)
        return False


def main() -> int:
    script_dir = os.path.dirname(os.path.abspath(__file__))
    default_input = os.path.join(script_dir, "input")
    default_output = os.path.join(script_dir, "output", "发票汇总.pptx")

    parser = argparse.ArgumentParser(description="将发票图片贴到PPT，每张图片一张幻灯片")
    parser.add_argument("--input-root", default=default_input, help="包含人员子文件夹的根目录")
    parser.add_argument("--persons", nargs="*", help="要处理的人员子文件夹名称（不提供则处理所有子文件夹）")
    parser.add_argument("--output", default=default_output, help="输出PPT文件路径")
    parser.add_argument("--pdf", action="store_true", help="同时将生成的PPT转换为PDF（需安装 LibreOffice）")

    args = parser.parse_args()
    if not os.path.isdir(args.input_root):
        print(f"输入目录不存在: {args.input_root}", file=sys.stderr)
        return 1
    os.makedirs(os.path.dirname(args.output), exist_ok=True)

    if args.persons:
        persons = args.persons
    else:
        persons = [d for d in sorted(os.listdir(args.input_root)) if os.path.isdir(os.path.join(args.input_root, d))]

    if not persons:
        print(f"输入目录下没有可处理的人员子目录: {args.input_root}", file=sys.stderr)
        return 1

    slide_count = create_ppt(args.input_root, persons, args.output)
    if slide_count == 0:
        print("未生成任何幻灯片，请检查输入目录结构或图片格式", file=sys.stderr)
        return 1

    if args.pdf:
        if not convert_to_pdf(args.output, os.path.dirname(args.output)):
            return 1

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
