#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""贴图版式引擎：按图片方向自动选择每页张数，输出 A4 幻灯片。

与 `create_invoice_ppt.py`（一图一页）的区别在于本模块会先看图片的宽高比：
竖版手机截图受限于高度，横排三列不会让单张变小，因此排成 3×2；横版发票受限
于宽度，上下两张同样不会变小，因此排成 1×2。夹在横版里的竖版图会单独占满一页，
否则几十行商品明细会糊成一片。
"""

from __future__ import annotations

import argparse
import csv
import math
import os
import statistics
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}

PAGE_SIZES = {"A4": (8.27, 11.69), "A4-landscape": (11.69, 8.27), "4:3": (10.0, 7.5)}

MARGIN = 0.40
HEADER = 0.72
FOOTER = 0.30
#: 宽高比低于此值视为「竖版手机截图」，整册按 3×2 网格排布。
TALL_ASPECT = 0.80
MANIFEST_ROWS_PER_PAGE = 34


def is_image_file(path: Path) -> bool:
    return path.is_file() and not path.name.startswith(".") and path.suffix.lower() in IMAGE_EXTENSIONS


def list_images(folder: str | Path) -> list[Path]:
    folder = Path(folder)
    return sorted(p for p in folder.iterdir() if is_image_file(p)) if folder.is_dir() else []


def aspect(path: str | Path) -> float:
    with Image.open(path) as im:
        return im.width / im.height


def _set_font(run, size: float, bold: bool = False, color: RGBColor | None = None, font: str = "Heiti SC") -> None:
    """同时设置拉丁与东亚字体，否则中文会回退到默认字体。"""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    if color is not None:
        run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.append(node)
        node.set("typeface", font)


def _textbox(slide, left, top, width, height, lines, align=PP_ALIGN.CENTER, font="Heiti SC"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for i, (text, size, bold, color) in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = align
        run = para.add_run()
        run.text = text
        _set_font(run, size, bold, color, font)
    return box


class SheetDeck:
    """A4 贴图册。"""

    def __init__(self, title: str, page: str = "A4", font: str = "Heiti SC") -> None:
        self.page_w, self.page_h = PAGE_SIZES[page]
        self.title = title
        self.font = font
        self.prs = Presentation()
        self.prs.slide_width = Inches(self.page_w)
        self.prs.slide_height = Inches(self.page_h)

    @property
    def usable(self) -> tuple[float, float]:
        return self.page_w - 2 * MARGIN, self.page_h - MARGIN - HEADER - FOOTER

    def _new_slide(self, heading: str):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        width = self.page_w - 2 * MARGIN
        _textbox(slide, MARGIN, 0.22, width, 0.44, [(heading, 13, True, RGBColor(0, 0, 0))], PP_ALIGN.LEFT, self.font)
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(HEADER - 0.08), Inches(width), Emu(6350))
        rule.fill.solid()
        rule.fill.fore_color.rgb = RGBColor(0x33, 0x33, 0x33)
        rule.line.fill.background()
        rule.shadow.inherit = False
        return slide

    def _place(self, slide, image: Path, caption: list[str], cx: float, cy: float, cw: float, ch: float, cap_h: float) -> None:
        max_w, max_h = cw * 0.95, ch - cap_h
        ratio = aspect(image)
        height = max_h
        width = height * ratio
        if width > max_w:
            width = max_w
            height = width / ratio
        x = cx + (cw - width) / 2
        y = cy + (max_h - height) / 2
        slide.shapes.add_picture(str(image), Inches(x), Inches(y), width=Inches(width), height=Inches(height))
        border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
        border.fill.background()
        border.line.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
        border.line.width = Pt(0.5)
        border.shadow.inherit = False
        if caption:
            lines = [(caption[0], 7.5, True, RGBColor(0, 0, 0))]
            lines += [(t, 6.5, False, RGBColor(0x44, 0x44, 0x44)) for t in caption[1:] if t]
            _textbox(slide, cx, cy + max_h + 0.03, cw, cap_h, lines, PP_ALIGN.CENTER, self.font)

    def add_grid(self, images: list[Path], captions: dict[str, list[str]], cols: int, rows: int, cap_h: float = 0.42) -> int:
        uw, uh = self.usable
        cw, ch = uw / cols, uh / rows
        per = cols * rows
        pages = math.ceil(len(images) / per) or 1
        for page in range(pages):
            slide = self._new_slide(f"{self.title}　|　第 {page + 1} / {pages} 页")
            for slot in range(per):
                i = page * per + slot
                if i >= len(images):
                    break
                r, c = divmod(slot, cols)
                self._place(slide, images[i], captions.get(images[i].name, []), MARGIN + c * cw, HEADER + r * ch, cw, ch, cap_h)
        return len(images)

    def add_pairs(self, images: list[Path], captions: dict[str, list[str]], cap_h: float = 0.34) -> int:
        """横版两张一页；竖版独占整页，避免细小明细被压扁。"""
        pages: list[list[Path]] = []
        current: list[Path] = []
        for image in images:
            if aspect(image) < 1.0:
                if current:
                    pages.append(current)
                    current = []
                pages.append([image])
            else:
                current.append(image)
                if len(current) == 2:
                    pages.append(current)
                    current = []
        if current:
            pages.append(current)

        uw, uh = self.usable
        for n, group in enumerate(pages, 1):
            slide = self._new_slide(f"{self.title}　|　第 {n} / {len(pages)} 页")
            portrait = len(group) == 1 and aspect(group[0]) < 1.0
            ch = uh if portrait else uh / 2
            for r, image in enumerate(group):
                self._place(slide, image, captions.get(image.name, []), MARGIN, HEADER + r * ch, uw, ch, cap_h)
        return len(images)

    def add_manifest(self, headers: list[str], widths: list[float], rows: list[list[str]], footer: list[str] | None = None) -> None:
        """在册子末尾追加清单页，便于核对总笔数与合计。"""
        uw = self.page_w - 2 * MARGIN
        scale = min(1.0, uw / sum(widths))
        widths = [w * scale for w in widths]
        pages = math.ceil(len(rows) / MANIFEST_ROWS_PER_PAGE) or 1
        for page in range(pages):
            slide = self._new_slide(f"{self.title}　|　凭证清单 {page + 1}/{pages}")
            chunk = rows[page * MANIFEST_ROWS_PER_PAGE : (page + 1) * MANIFEST_ROWS_PER_PAGE]
            last = page == pages - 1
            n_rows = len(chunk) + 1 + (1 if footer and last else 0)
            table = slide.shapes.add_table(
                n_rows, len(headers), Inches(MARGIN), Inches(HEADER + 0.05), Inches(sum(widths)), Inches(0.26 * n_rows)
            ).table
            for i, w in enumerate(widths):
                table.columns[i].width = Inches(w)
            for i, text in enumerate(headers):
                self._fill_cell(table.cell(0, i), text, 8.5, True, PP_ALIGN.CENTER)
            for r, row in enumerate(chunk, 1):
                for c, text in enumerate(row):
                    self._fill_cell(table.cell(r, c), text, 8, False, PP_ALIGN.LEFT if c == 2 else PP_ALIGN.CENTER)
            if footer and last:
                for c, text in enumerate(footer):
                    self._fill_cell(table.cell(n_rows - 1, c), text, 9, True, PP_ALIGN.CENTER)
            table.rows[0].height = Inches(0.26)

    def _fill_cell(self, cell, text: str, size: float, bold: bool, align) -> None:
        cell.text = text
        for para in cell.text_frame.paragraphs:
            para.alignment = align
            for run in para.runs:
                _set_font(run, size, bold, None, self.font)

    def save(self, path: str | Path) -> None:
        self.prs.save(str(path))


def choose_layout(images: list[Path]) -> str:
    """按宽高比中位数决定版式：竖版图走网格，横版图走上下配对。"""
    if not images:
        return "grid"
    return "grid" if statistics.median(aspect(p) for p in images) < TALL_ASPECT else "pairs"


def build_deck(
    images: list[Path],
    output: str | Path,
    title: str,
    captions: dict[str, list[str]] | None = None,
    manifest: dict | None = None,
    layout: str = "auto",
    page: str = "A4",
    grid: tuple[int, int] = (3, 2),
) -> tuple[int, int]:
    """返回 (幻灯片数, 已贴图片数)。"""
    captions = captions or {}
    deck = SheetDeck(title, page=page)
    mode = choose_layout(images) if layout == "auto" else layout
    placed = deck.add_grid(images, captions, *grid) if mode == "grid" else deck.add_pairs(images, captions)
    if manifest:
        deck.add_manifest(manifest["headers"], manifest["widths"], manifest["rows"], manifest.get("footer"))
    deck.save(output)
    return len(deck.prs.slides._sldIdLst), placed


def load_captions(path: str | Path) -> dict[str, list[str]]:
    """读取标注 CSV：第一列为文件名，其余列依次作为该图下方的标注行。"""
    result: dict[str, list[str]] = {}
    with open(path, encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle):
            if len(row) >= 2 and row[0].strip():
                result[row[0].strip()] = [c.strip() for c in row[1:] if c.strip()]
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="把一个文件夹的凭证图片贴成 A4 幻灯片（自动按方向排版）")
    parser.add_argument("folder", help="包含图片的文件夹")
    parser.add_argument("--output", required=True, help="输出 PPTX 路径")
    parser.add_argument("--title", default="", help="页眉标题，默认用文件夹名")
    parser.add_argument("--captions", help="标注 CSV：文件名,第一行,第二行,...")
    parser.add_argument("--layout", choices=["auto", "grid", "pairs"], default="auto")
    parser.add_argument("--page", choices=sorted(PAGE_SIZES), default="A4")
    parser.add_argument("--grid", default="3x2", help="grid 版式的列x行，默认 3x2")
    args = parser.parse_args()

    images = list_images(args.folder)
    if not images:
        print(f"没有找到图片: {args.folder}")
        return 1

    cols, rows = (int(v) for v in args.grid.lower().split("x"))
    captions = load_captions(args.captions) if args.captions else {}
    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
    slides, placed = build_deck(
        images,
        args.output,
        args.title or Path(args.folder).name,
        captions=captions,
        layout=args.layout,
        page=args.page,
        grid=(cols, rows),
    )
    print(f"✓ {placed} 张图片 → {slides} 页: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
